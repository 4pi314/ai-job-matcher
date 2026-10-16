from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from document_parser import extract_text_from_file


class NamedBytesIO(BytesIO):
    """An in-memory file that also has a filename."""

    def __init__(self, content: bytes, name: str):
        super().__init__(content)
        self.name = name


def test_extracts_text_from_txt():
    uploaded_file = NamedBytesIO(
        b"Python developer with SQL experience.",
        "resume.txt",
    )

    result = extract_text_from_file(uploaded_file)

    assert result == "Python developer with SQL experience."


def test_extracts_text_from_docx():
    document = Document()
    document.add_paragraph("Python developer")
    document.add_paragraph("Experienced with Docker and Git")

    file_buffer = BytesIO()
    document.save(file_buffer)

    uploaded_file = NamedBytesIO(
        file_buffer.getvalue(),
        "resume.docx",
    )

    result = extract_text_from_file(uploaded_file)

    assert "Python developer" in result
    assert "Docker and Git" in result


def test_reads_valid_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)

    file_buffer = BytesIO()
    writer.write(file_buffer)

    uploaded_file = NamedBytesIO(
        file_buffer.getvalue(),
        "resume.pdf",
    )

    result = extract_text_from_file(uploaded_file)

    assert result == ""


def test_rejects_unsupported_file_type():
    uploaded_file = NamedBytesIO(
        b"image content",
        "resume.jpg",
    )

    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text_from_file(uploaded_file)